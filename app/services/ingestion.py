from langchain_core.documents import Document
from pathlib import Path
from typing import Iterable
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument

SUPPORTED_FILE_TYPES = [".pdf", ".txt",".md",".docx"]

def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        loader = PyPDFLoader(str(path))
        return loader.load()
    elif suffix == ".txt" or suffix == ".md":
        loader = TextLoader(str(path))
        return loader.load()
    elif suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join(para.text for para in doc.paragraphs if para.text.strip())
        return [Document(page_content=text, metadata={"source": str(path)})]
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size = 900, chunk_overlap=120, add_start_index=True)
    return splitter.split_documents(list(docs))